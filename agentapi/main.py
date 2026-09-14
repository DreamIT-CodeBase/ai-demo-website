from fastapi import FastAPI, File, Form, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, Response
from azure.ai.projects import AIProjectClient
from azure.identity import DefaultAzureCredential
import json
import re
from pptx import Presentation
import os
import tempfile
import threading
import base64

app = FastAPI(title="AI Agent API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

PROJECT_ENDPOINT = "https://raghavk-8931-resource.services.ai.azure.com/api/projects/raghavk-8931"

project = AIProjectClient(
    endpoint=PROJECT_ENDPOINT,
    credential=DefaultAzureCredential(),
    allow_preview=True
)


DOCUMENT_EXTENSIONS = {".pdf", ".doc", ".docx", ".pptx", ".txt", ".json", ".md", ".html"}
TABULAR_EXTENSIONS = {".csv", ".xlsx", ".xls"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
ALLOWED_EXTENSIONS = DOCUMENT_EXTENSIONS | TABULAR_EXTENSIONS | IMAGE_EXTENSIONS

IMAGE_MIME_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

def classify_extension(extension: str) -> str | None:
    if extension in DOCUMENT_EXTENSIONS:
        return "document"
    if extension in TABULAR_EXTENSIONS:
        return "tabular"
    if extension in IMAGE_EXTENSIONS:
        return "image"
    return None

# In-memory per-conversation file state, so the client never has to remember
# or resend file/vector-store IDs between turns of the same conversation.
#
# Maps conversation_id -> {
#   "vector_store_id": str | None,           - one growing store; every uploaded
#                                               document gets added into it, so
#                                               earlier documents stay searchable
#                                               alongside new ones
#   "code_interpreter_file_ids": list[str],  - accumulates, most recent
#                                               CODE_INTERPRETER_SLOT_COUNT kept
#   "uploaded_filenames": list[str],         - every filename uploaded so far,
#                                               used to nudge the model to look
#                                               at a freshly uploaded file
# }
#
# NOTE: this is process-local. It resets on restart and won't be shared
# across multiple backend instances behind a load balancer. For production,
# swap this dict for Redis or a small database table keyed by conversation_id.
CONVERSATION_FILES: dict[str, dict] = {}
CONVERSATION_FILES_LOCK = threading.Lock()

# Latest generated Code Interpreter file metadata by filename.
# Used only as a fallback when a frontend request contains Foundry's
# sandbox:/mnt/data/<filename> link but the file event was not retained client-side.
GENERATED_FILES_BY_FILENAME: dict[str, dict] = {}
GENERATED_FILES_LOCK = threading.Lock()

# How many spreadsheet files Code Interpreter can hold onto at once for a
# conversation. Must match the number of {{code_interpreter_file_id_N}}
# slots configured on the agent (see setup_code_interpreter.py). Oldest
# tabular file is dropped once this limit is exceeded.
CODE_INTERPRETER_SLOT_COUNT = 3

def get_conversation_files(conversation_id: str | None) -> dict:
    empty = {
        "vector_store_id": None,
        "code_interpreter_file_ids": [],
        "uploaded_filenames": [],
    }

    if not conversation_id:
        return dict(empty)

    with CONVERSATION_FILES_LOCK:
        stored = CONVERSATION_FILES.get(conversation_id, empty)
        # Merge with `empty` so conversations created before this change
        # (old shape: single code_interpreter_file_id) don't KeyError.
        return {**empty, **stored}

def set_conversation_files(conversation_id: str, files: dict):
    with CONVERSATION_FILES_LOCK:
        CONVERSATION_FILES[conversation_id] = files

KNOWLEDGE_TOOL_TYPES = {
    "fabric_iq_preview",
    "work_iq_preview",
    "file_search",
    "web_search",
    "azure_ai_search",
    "bing_grounding",
    "bing_custom_search"
}

def send_event(event_type, data):
    return json.dumps({"event": event_type, **data}) + "\n"

# ---------------------------------------------------------------------------
# Citation artifact cleanup
#
# GPT-5-family models occasionally leak their internal file-search citation
# reference tokens into the actual output text instead of converting them
# to real citations. This shows up in two different forms:
#
#   1. Visible bracket form:   "...retail data.【filecite】turn0file7【filecite】turn0file14】..."
#   2. Invisible-marker form:  "...650-hour cap.<PUA><PUA><PUA>turn0file15<PUA>..."
#      (renders as empty boxes/tofu glyphs, since those are Unicode Private
#      Use Area control characters with no real glyph - browsers show a
#      placeholder box for them)
#
# Both carry no useful information for the end user, so strip them out of
# every streamed text delta before it reaches the browser.
# ---------------------------------------------------------------------------
CITATION_ARTIFACT_RE = re.compile(
    r"[【〔\[]?\s*filecite\s*[】〕\]]?"       # the "filecite" label, with or without brackets
    r"|turn\d+file\d+(?:[A-Za-z]\d+-[A-Za-z]?\d*)?"  # e.g. turn0file7, turn0file14L1-L3
    r"|[【】〔〕]"                             # any leftover stray bracket characters
    r"|[\uE000-\uF8FF]"                      # Unicode Private Use Area (BMP)
    r"|[\U000F0000-\U000FFFFD]"              # Private Use Area, plane 15
    r"|[\U00100000-\U0010FFFD]"              # Private Use Area, plane 16
)

def strip_citation_artifacts(text):
    if not text:
        return text
    return CITATION_ARTIFACT_RE.sub("", text)

# ---------------------------------------------------------------------------
# PowerPoint text extraction
#
# Azure's vector store ingestion for .pptx is unreliable - uploads can
# succeed with no error, but the slide content never actually becomes
# searchable, so File Search returns nothing for it. To avoid depending on
# Azure's native Office-format parsing, we extract the text ourselves
# (slide text, tables, and speaker notes) and upload that as a plain .txt
# file instead, which Azure indexes reliably.
# ---------------------------------------------------------------------------
def extract_pptx_text(path: str) -> str:
    presentation = Presentation(path)
    slide_sections = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        lines = [f"--- Slide {slide_index} ---"]

        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        lines.append(text)

            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))

        if slide.has_notes_slide:
            notes_text = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes_text:
                lines.append(f"[Speaker notes] {notes_text}")

        slide_sections.append("\n".join(lines))

    return "\n\n".join(slide_sections)

# ---------------------------------------------------------------------------
# JSON / plain text normalization
#
# Two more formats were also failing to actually become searchable in
# Azure's vector store, the same way .pptx did:
#   - .json: Azure appears to special-case the .json extension (likely
#     expecting a specific schema) and silently fails to index arbitrary
#     JSON content. Re-serializing as pretty-printed text and uploading it
#     with a .txt extension avoids that special-casing, and is also better
#     for search relevance than dense, unindented JSON.
#   - .txt / .md: re-saved through the same guaranteed-clean UTF-8 .txt
#     pipeline as everything else here, removing any dependency on how
#     Azure's parser dispatches on the original extension.
# ---------------------------------------------------------------------------
def normalize_json_text(raw_bytes: bytes) -> str:
    try:
        parsed = json.loads(raw_bytes.decode("utf-8"))
        return json.dumps(parsed, indent=2, ensure_ascii=False)
    except (ValueError, UnicodeDecodeError):
        # Not valid JSON, or not UTF-8 - fall back to whatever text we can
        # decode rather than failing the upload outright.
        return raw_bytes.decode("utf-8", errors="replace")

def read_text_file(raw_bytes: bytes) -> str:
    return raw_bytes.decode("utf-8", errors="replace")

def get_annotation_field(source, *names):
    """Try several possible attribute/key names on an annotation-like
    object, and also look one level deeper under a
    'container_file_citation' sub-object, since some SDK/API versions
    nest type-specific fields there instead of putting them directly on
    the annotation."""
    if source is None:
        return None

    for name in names:
        value = getattr(source, name, None)
        if value is None and isinstance(source, dict):
            value = source.get(name)
        if value is not None:
            return value

    nested = getattr(source, "container_file_citation", None)
    if nested is None and isinstance(source, dict):
        nested = source.get("container_file_citation")
    if nested is not None:
        for name in names:
            value = getattr(nested, name, None)
            if value is None and isinstance(nested, dict):
                value = nested.get(name)
            if value is not None:
                return value

    return None

def serialize_tool(tool):
    if hasattr(tool, "as_dict"):
        data = tool.as_dict()
    elif isinstance(tool, dict):
        data = tool
    else:
        data = {}

    if not data:
        data = {
            "type": getattr(tool, "type", None),
            "name": getattr(tool, "name", None)
        }

    return data

def get_tool_type(tool):
    return str(tool.get("type") or "").lower()

def normalize_tools(tools):
    normalized = []

    for tool in tools or []:
        data = serialize_tool(tool)

        if not data.get("type"):
            continue

        normalized.append(data)

    return normalized

@app.get("/health")
def health():
    return {"status": "ok"}

@app.get("/agents")
def list_agents():
    try:
        return [
            {
                "name": getattr(agent, "name", None),
                "description": getattr(agent, "description", None),
                "kind": getattr(agent, "kind", None)
            }
            for agent in project.agents.list(limit=100)
        ]
    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to list agents: {str(error)}"
        )

@app.get("/agents/{agent_name}/metadata")
def get_agent_metadata(agent_name: str):
    try:
        versions = list(
            project.agents.list_versions(
                agent_name=agent_name,
                limit=100,
                order="desc"
            )
        )

        if not versions:
            raise HTTPException(
                status_code=404,
                detail=f"No versions found for agent '{agent_name}'"
            )

        latest = versions[0]

        version = getattr(latest, "version", None)

        if version is None and isinstance(latest, dict):
            version = latest.get("version")

        if version is None:
            raise HTTPException(
                status_code=500,
                detail=f"Could not determine version for agent '{agent_name}'"
            )

        agent_version = project.agents.get_version(
            agent_name=agent_name,
            agent_version=version
        )

        definition = getattr(agent_version, "definition", None)

        if definition is None and isinstance(agent_version, dict):
            definition = agent_version.get("definition")

        if definition is None:
            raise HTTPException(
                status_code=404,
                detail=f"No definition found for agent '{agent_name}' version '{version}'"
            )

        if hasattr(definition, "as_dict"):
            definition_data = definition.as_dict()
        elif isinstance(definition, dict):
            definition_data = definition
        else:
            definition_data = {}

        model = getattr(definition, "model", None)

        if model is None:
            model = definition_data.get("model")

        raw_tools = getattr(definition, "tools", None)

        if raw_tools is None:
            raw_tools = definition_data.get("tools", [])

        tools = normalize_tools(raw_tools)

        knowledge_sources = [
            tool
            for tool in tools
            if get_tool_type(tool) in KNOWLEDGE_TOOL_TYPES
        ]

        action_tools = [
            tool
            for tool in tools
            if get_tool_type(tool) not in KNOWLEDGE_TOOL_TYPES
        ]

        return {
            "name": agent_name,
            "version": version,
            "description": getattr(agent_version, "description", None),
            "model": model,
            "tools": tools,
            "knowledge_sources": knowledge_sources,
            "action_tools": action_tools,
            "tool_counts": {
                "total": len(tools),
                "knowledge_sources": len(knowledge_sources),
                "action_tools": len(action_tools),
                "openapi": sum(
                    1 for tool in tools
                    if get_tool_type(tool) == "openapi"
                ),
                "mcp": sum(
                    1 for tool in tools
                    if get_tool_type(tool) == "mcp"
                ),
                "fabric_iq": sum(
                    1 for tool in tools
                    if get_tool_type(tool) == "fabric_iq_preview"
                ),
                "work_iq": sum(
                    1 for tool in tools
                    if get_tool_type(tool) == "work_iq_preview"
                ),
                "file_search": sum(
                    1 for tool in tools
                    if get_tool_type(tool) == "file_search"
                ),
                "web_search": sum(
                    1 for tool in tools
                    if get_tool_type(tool) == "web_search"
                )
            },
            "definition": definition_data
        }

    except HTTPException:
        raise

    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to retrieve agent metadata: {str(error)}"
        )

@app.get("/files/by-filename")
def download_container_file_by_filename(filename: str):
    key = os.path.basename(filename.strip())
    with GENERATED_FILES_LOCK:
        file_info = GENERATED_FILES_BY_FILENAME.get(key)

    if not file_info:
        raise HTTPException(status_code=404, detail=f"Generated file not found: {key}")

    return download_container_file(
        file_info["container_id"],
        file_info["file_id"],
        key,
    )


@app.get("/files/{container_id}/{file_id}")
def download_container_file(container_id: str, file_id: str, filename: str = "generated-file"):
    """Serves a file the model generated via Code Interpreter (e.g. a
    document the user asked the agent to create). Azure/OpenAI cite these
    generated files as 'container_file_citation' annotations, which we
    detect in the /chat stream below and turn into a download link
    pointing back at this endpoint."""
    try:
        openai_files = project.get_openai_client()
        result = openai_files.containers.files.content.retrieve(
            container_id=container_id,
            file_id=file_id,
        )

        if hasattr(result, "read"):
            file_bytes = result.read()
        elif isinstance(result, (bytes, bytearray)):
            file_bytes = bytes(result)
        else:
            file_bytes = result.content

    except Exception as error:
        raise HTTPException(
            status_code=502,
            detail=f"Failed to retrieve generated file from Azure: {str(error)}"
        )

    return Response(
        content=file_bytes,
        media_type="application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@app.post("/chat")
async def chat(
    agent: str = Form(...),
    message: str = Form(""),
    conversation_id: str | None = Form(None),
    vector_store_id: str | None = Form(None),
    files: list[UploadFile] = File(default_factory=list)
):
    async def generate():
        temp_paths = []

        try:
            openai = project.get_openai_client(agent_name=agent)
            openai_files = project.get_openai_client()

            if conversation_id:
                current_conversation_id = conversation_id
            else:
                conversation = openai.conversations.create()
                current_conversation_id = conversation.id

            conversation_files = get_conversation_files(current_conversation_id)

            # An explicit vector_store_id from the client (if the frontend ever
            # sends one) overrides what's stored server-side for this turn.
            if vector_store_id:
                conversation_files["vector_store_id"] = vector_store_id

            current_vector_store_id = conversation_files.get("vector_store_id")
            current_code_interpreter_file_ids = list(
                conversation_files.get("code_interpreter_file_ids", [])
            )
            uploaded_filenames = list(conversation_files.get("uploaded_filenames", []))

            just_uploaded_filenames = []
            image_inputs = []

            # FastAPI can hand back a single empty UploadFile (filename "")
            # when no files were actually attached - filter those out.
            incoming_files = [f for f in (files or []) if f and f.filename]

            for incoming_file in incoming_files:
                extension = os.path.splitext(incoming_file.filename or "")[1].lower()
                file_kind = classify_extension(extension)

                if file_kind is None:
                    yield send_event("error", {
                        "message": f"Unsupported file type: {extension or 'unknown'} ({incoming_file.filename})"
                    })
                    return

                file_bytes = await incoming_file.read()

                with tempfile.NamedTemporaryFile(
                    delete=False,
                    suffix=extension
                ) as temp_file:
                    temp_file.write(file_bytes)
                    temp_path = temp_file.name
                    temp_paths.append(temp_path)

                yield send_event("activity", {
                    "type": "file_upload",
                    "status": "in_progress",
                    "name": f"Uploading {incoming_file.filename}"
                })

                if file_kind == "document":
                    # Reuse the same vector store for the whole conversation so
                    # every document uploaded so far stays searchable together,
                    # instead of each new upload replacing the last one.
                    if not current_vector_store_id:
                        vector_store = openai_files.vector_stores.create(
                            name=f"portal-{agent}-{current_conversation_id}"
                        )
                        current_vector_store_id = vector_store.id

                    upload_path = temp_path
                    extracted_text = None

                    if extension == ".pptx":
                        # Azure's vector store often accepts a .pptx upload
                        # with no error, but never actually indexes its
                        # content, so file_search silently finds nothing.
                        # Extract the text ourselves and upload plain text
                        # instead, which Azure indexes reliably.
                        extracted_text = extract_pptx_text(temp_path)
                        if not extracted_text.strip():
                            extracted_text = (
                                f"(The PowerPoint file '{incoming_file.filename}' "
                                "had no extractable text - it may be image-only.)"
                            )

                    elif extension == ".json":
                        extracted_text = normalize_json_text(file_bytes)

                    elif extension in (".txt", ".md"):
                        extracted_text = read_text_file(file_bytes)

                    if extracted_text is not None:
                        safe_base = re.sub(
                            r"[^A-Za-z0-9_.-]", "_",
                            os.path.splitext(incoming_file.filename or "document")[0]
                        )[:60]

                        with tempfile.NamedTemporaryFile(
                            delete=False,
                            mode="w",
                            encoding="utf-8",
                            suffix=f"__{safe_base}.txt"
                        ) as txt_file:
                            txt_file.write(extracted_text)
                            upload_path = txt_file.name
                            temp_paths.append(upload_path)

                    with open(upload_path, "rb") as file_handle:
                        openai_files.vector_stores.files.upload_and_poll(
                            vector_store_id=current_vector_store_id,
                            file=file_handle
                        )

                elif file_kind == "tabular":
                    with open(temp_path, "rb") as file_handle:
                        uploaded_file = openai_files.files.create(
                            purpose="assistants",
                            file=file_handle
                        )

                    current_code_interpreter_file_ids.append(uploaded_file.id)
                    # Keep only the most recent N tabular files - matches the
                    # number of {{code_interpreter_file_id_N}} slots the agent
                    # was configured with.
                    current_code_interpreter_file_ids = current_code_interpreter_file_ids[
                        -CODE_INTERPRETER_SLOT_COUNT:
                    ]

                else:  # file_kind == "image"
                    # Images don't need file_search or code_interpreter - they
                    # go straight into this turn's input as vision content, the
                    # same way the standard Responses API handles image input.
                    mime_type = IMAGE_MIME_TYPES.get(extension, "application/octet-stream")
                    encoded = base64.b64encode(file_bytes).decode("utf-8")
                    image_inputs.append({
                        "type": "input_image",
                        "image_url": f"data:{mime_type};base64,{encoded}"
                    })

                just_uploaded_filenames.append(incoming_file.filename)
                uploaded_filenames.append(incoming_file.filename)

                yield send_event("activity", {
                    "type": "file_upload",
                    "status": "completed",
                    "name": f"{incoming_file.filename} ready"
                })

            if incoming_files:
                conversation_files["vector_store_id"] = current_vector_store_id
                conversation_files["code_interpreter_file_ids"] = current_code_interpreter_file_ids
                conversation_files["uploaded_filenames"] = uploaded_filenames
                set_conversation_files(current_conversation_id, conversation_files)

            structured_inputs = {}

            if current_vector_store_id:
                structured_inputs["vector_store_id"] = current_vector_store_id

            for index in range(CODE_INTERPRETER_SLOT_COUNT):
                slot_name = f"code_interpreter_file_id_{index + 1}"
                if index < len(current_code_interpreter_file_ids):
                    structured_inputs[slot_name] = current_code_interpreter_file_ids[index]

            extra_body = {
                "structured_inputs": structured_inputs
            } if structured_inputs else None

            effective_message = message or "Please analyze the uploaded document."

            # Nudge the model to actually search/inspect the file(s) that were
            # just uploaded, instead of answering from earlier conversation
            # context as if nothing changed. Also reminds it that other
            # previously uploaded files are still available.
            if just_uploaded_filenames:
                other_files = [
                    name for name in uploaded_filenames if name not in just_uploaded_filenames
                ]
                if len(just_uploaded_filenames) == 1:
                    note = f"[A new file, '{just_uploaded_filenames[0]}', was just uploaded."
                else:
                    note = (
                        "[" + str(len(just_uploaded_filenames)) + " new files were just uploaded: "
                        + ", ".join(just_uploaded_filenames) + "."
                    )
                if other_files:
                    note += (
                        " Previously uploaded files are still available too: "
                        + ", ".join(other_files) + "."
                    )
                note += " Make sure to search/analyze the newly uploaded file(s) for this question.]"
                effective_message = f"{note}\n\n{effective_message}"

            print("Vector Store:", current_vector_store_id)
            print("Code Interpreter Files:", current_code_interpreter_file_ids)
            print("Structured Inputs:", structured_inputs)
            print("Extra Body:", extra_body)
            print("Image inputs attached this turn:", len(image_inputs))

            if image_inputs:
                # Images ride along as vision content on this turn's message,
                # instead of the plain string input used otherwise.
                input_payload = [{
                    "role": "user",
                    "content": [{"type": "input_text", "text": effective_message}] + image_inputs
                }]
            else:
                input_payload = effective_message

            response_stream = openai.responses.create(
                conversation=current_conversation_id,
                input=input_payload,
                stream=True,
                extra_body=extra_body
            )

            emitted_files = set()

            for event in response_stream:
                event_type = getattr(event, "type", None)

                if event_type in {
                    "response.output_item.added",
                    "response.output_item.done"
                }:
                    item = getattr(event, "item", None)
                    print("ITEM TYPE:", getattr(item, "type", None))
                    print("ITEM:", item)

                if not event_type or event_type == "keepalive":
                    continue

                print("EVENT:", event_type)

                if event_type == "response.reasoning_summary_text.delta":
                    delta = getattr(event, "delta", None)

                    if delta:
                        yield send_event("reasoning", {
                            "delta": delta
                        })

                    continue

                if event_type == "response.output_item.added":
                    item = getattr(event, "item", None)

                    if item:
                        item_type = getattr(item, "type", None)

                        if item_type == "openapi_call":
                            yield send_event("activity", {
                                "type": "openapi_call",
                                "status": "in_progress",
                                "name": getattr(item, "name", None)
                            })

                        elif item_type == "openapi_call_output":
                            yield send_event("activity", {
                                "type": "openapi_call_output",
                                "status": "in_progress"
                            })

                    continue

                if event_type == "response.output_item.done":
                    item = getattr(event, "item", None)

                    if item:
                        item_type = getattr(item, "type", None)

                        if item_type == "openapi_call":
                            yield send_event("activity", {
                                "type": "openapi_call",
                                "status": "completed",
                                "name": getattr(item, "name", None)
                            })

                        elif item_type == "openapi_call_output":
                            yield send_event("activity", {
                                "type": "openapi_call_output",
                                "status": "completed"
                            })

                    continue

                if event_type == "response.output_text.delta":
                    delta = getattr(event, "delta", None)

                    if delta:
                        cleaned_delta = strip_citation_artifacts(delta)

                        if cleaned_delta:
                            yield send_event("text", {
                                "delta": cleaned_delta
                            })

                    continue

                # Catch-all diagnostic: log the full shape of ANY event whose
                # type mentions "annotation", so we can see exactly what
                # Azure is sending if the specific handler below doesn't
                # match it. Safe to leave in - it's a no-op for every other
                # event type.
                if "annotation" in (event_type or "").lower():
                    try:
                        raw_shape = event.model_dump() if hasattr(event, "model_dump") else str(event)
                    except Exception:
                        raw_shape = repr(event)
                    print(f"ANNOTATION-RELATED EVENT [{event_type}]:", raw_shape)

                if event_type == "response.output_text.annotation.added":
                    annotation = getattr(event, "annotation", None)
                    ann_type = get_annotation_field(annotation, "type")

                    if ann_type == "container_file_citation" or (
                        ann_type is None and get_annotation_field(annotation, "container_id")
                    ):
                        container_id = get_annotation_field(annotation, "container_id")
                        file_id = get_annotation_field(annotation, "file_id", "id")
                        filename = get_annotation_field(annotation, "filename", "file_name") or "generated-file"

                        print("GENERATED FILE ANNOTATION:", container_id, file_id, filename)

                        if container_id and file_id and (container_id, file_id) not in emitted_files:
                            emitted_files.add((container_id, file_id))
                            with GENERATED_FILES_LOCK:
                                GENERATED_FILES_BY_FILENAME[filename] = {
                                    "container_id": container_id,
                                    "file_id": file_id,
                                }
                            yield send_event("file", {
                                "container_id": container_id,
                                "file_id": file_id,
                                "filename": filename
                            })
                        elif not (container_id and file_id):
                            print(
                                "  WARNING: annotation matched but container_id/file_id "
                                "still missing - check the raw shape logged above."
                            )

                    continue

                if event_type in {
                    "response.mcp_call.in_progress",
                    "response.mcp_call.completed",
                    "response.mcp_call.failed",
                    "response.mcp_list_tools.in_progress",
                    "response.mcp_list_tools.completed",
                    "response.mcp_list_tools.failed",
                    "response.file_search_call.in_progress",
                    "response.file_search_call.completed",
                    "response.file_search_call.failed",
                    "response.code_interpreter_call.in_progress",
                    "response.code_interpreter_call.interpreting",
                    "response.code_interpreter_call.completed",
                    "response.code_interpreter_call.failed"
                }:
                    activity_data = {
                        "type": event_type,
                        "status": (
                            "completed"
                            if event_type.endswith(".completed")
                            else "failed"
                            if event_type.endswith(".failed")
                            else "in_progress"
                        )
                    }

                    for attribute in [
                        "item_id",
                        "output_index",
                        "server_label",
                        "name"
                    ]:
                        value = getattr(event, attribute, None)

                        if value is not None:
                            activity_data[attribute] = value

                    yield send_event("activity", activity_data)
                    continue

                if event_type == "response.completed":
                    # Safety net: if Azure never streamed individual
                    # "annotation.added" events for a generated file (this
                    # can vary by API/SDK version), the citation may still
                    # be sitting in the final response object's output
                    # items. Scan those too, skipping anything already
                    # emitted above.
                    full_response = getattr(event, "response", None)
                    output_items = getattr(full_response, "output", None) or []

                    for output_item in output_items:
                        content_list = getattr(output_item, "content", None) or []

                        for content_item in content_list:
                            annotations = getattr(content_item, "annotations", None) or []

                            for annotation in annotations:
                                ann_type = get_annotation_field(annotation, "type")

                                if ann_type != "container_file_citation":
                                    continue

                                container_id = get_annotation_field(annotation, "container_id")
                                file_id = get_annotation_field(annotation, "file_id", "id")
                                filename = get_annotation_field(
                                    annotation, "filename", "file_name"
                                ) or "generated-file"

                                if (
                                    container_id and file_id
                                    and (container_id, file_id) not in emitted_files
                                ):
                                    emitted_files.add((container_id, file_id))
                                    print(
                                        "GENERATED FILE ANNOTATION (from completed response):",
                                        container_id, file_id, filename
                                    )
                                    yield send_event("file", {
                                        "container_id": container_id,
                                        "file_id": file_id,
                                        "filename": filename
                                    })

                    yield send_event("completed", {
                        "conversation_id": current_conversation_id,
                        "vector_store_id": current_vector_store_id,
                        "code_interpreter_file_ids": current_code_interpreter_file_ids,
                        "uploaded_filenames": uploaded_filenames
                    })

                    continue

        except GeneratorExit:
            return

        except Exception as error:
            import traceback
            traceback.print_exc()
            yield send_event("error", {
                "message": str(error),
                "error_type": type(error).__name__
            })

        finally:
            for temp_path in temp_paths:
                if temp_path and os.path.exists(temp_path):
                    os.remove(temp_path)

    return StreamingResponse(
        generate(),
        media_type="application/x-ndjson"
    )